"""Gera entregaveis/Match-Entregaveis.docx (Word, com imagens) a partir dos mesmos Markdown do PDF.

Como rodar (depois do `node gerar_pdf.mjs`, que renderiza os diagramas em docs/entrega/diagramas/):
    cd docs/entrega
    python gerar_docx.py
Precisa de: pandoc (https://pandoc.org) e python-docx (`pip install python-docx`).

O que faz: junta os 6 entregáveis na ordem do regulamento; troca cada bloco Mermaid pela imagem do diagrama já
renderizada; troca cada galeria de telas por uma tabela de 4 colunas (imagem + legenda); links relativos viram links
do GitHub. Não edita nenhum doc de origem.
"""

from __future__ import annotations

import html
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent.parent
DIAGRAMAS = AQUI / "diagramas"
SAIDA = RAIZ / "entregaveis" / "Match-Entregaveis.docx"
REPO_WEB = "https://github.com/scorpion-bits/Hackaton-Ifood/blob/main/"
VERMELHO = RGBColor(0xEA, 0x1D, 0x2C)

TITULO = "Match — pedido em grupo com agentes no MCP do iFood"
EVENTO = 'Hackathon SEMCOMP × iFood 2026 · "Vendendo com Agentes" · São Carlos'
EQUIPE = "Thales Miguel Hajes · Milan Bahrami · Zenon · Antonio"
DATA = "18 de outubro de 2026"
AVISO = "Protótipo de hackathon · não afiliado ao iFood. Dados simulados aparecem com o selo DEMONSTRAÇÃO."
TELAS_CAPA = ["docs/entrega/telas/06-pedido-do-grupo.png", "docs/entrega/telas/real-pedido.png",
              "docs/entrega/telas/10-confira.png"]


def ler(rel: str) -> str:
    return (RAIZ / rel).read_text(encoding="utf-8").replace("\r\n", "\n")


def tirar_h1(md: str) -> str:
    return re.sub(r"^\s*# [^\n]*\n", "", md, count=1)


def dividir_tecnico() -> tuple[str, str]:
    md = tirar_h1(ler("docs/tecnico.md"))
    m = re.search(r"^## 5\. ", md, flags=re.M)
    if not m:
        return md, md
    nota = ("> Continuação de `docs/tecnico.md` (seções 5 a 12). Visão geral, componentes, sequência e fluxo do motor "
            "estão no entregável 02.\n\n")
    return md[:m.start()], nota + md[m.start():]


TEC02, TEC05 = dividir_tecnico()
ENTREGAVEIS = [
    ("01", "Proposta de Negócio", "docs/proposta-negocio.md", tirar_h1(ler("docs/proposta-negocio.md")), "docs"),
    ("02", "Desenho de Solução", "docs/tecnico.md §1–4", TEC02, "docs"),
    ("03", "Racional da Experiência", "docs/racional-ux.md", tirar_h1(ler("docs/racional-ux.md")), "docs"),
    ("04", "Interface Visual", "docs/entrega/04-interface.md · docs/entrega/telas/",
     tirar_h1(ler("docs/entrega/04-interface.md")), "docs/entrega"),
    ("05", "Documento Técnico Resumido", "docs/tecnico.md §5–12", TEC05, "docs"),
    ("06", "Feedback sobre os MCPs", "docs/feedback-mcp.md", tirar_h1(ler("docs/feedback-mcp.md")), "docs"),
]

galerias: list[list[tuple[Path, str, str]]] = []
cont_diagrama = 0


def sem_tags(s: str) -> str:
    return html.unescape(re.sub(r"<[^>]+>", "", s)).strip()


def converter(md: str, base: str) -> str:
    global cont_diagrama

    def mermaid(_m: re.Match) -> str:
        global cont_diagrama
        cont_diagrama += 1
        png = DIAGRAMAS / f"diagrama-{cont_diagrama}.png"
        if not png.exists():
            return "\n> (diagrama no repositório: rode `node gerar_pdf.mjs` antes para incluí-lo aqui)\n"
        return f"\n![]({png.as_posix()})\n"

    md = re.sub(r"```mermaid\n.*?```", mermaid, md, flags=re.S)

    def galeria(m: re.Match) -> str:
        itens = []
        for f in re.finditer(r'<figure><img src="([^"]+)"[^>]*><figcaption>(.*?)</figcaption></figure>', m.group(0),
                             flags=re.S):
            leg = f.group(2)
            t = re.search(r"<b>(.*?)</b>", leg)
            titulo = sem_tags(t.group(1)) if t else ""
            texto = sem_tags(leg[t.end():] if t else leg)
            itens.append(((RAIZ / base / f.group(1)).resolve(), titulo, texto))
        galerias.append(itens)
        return f"\n\nGALERIA-{len(galerias) - 1}\n\n"

    md = re.sub(r'<div class="galeria">.*?</div>', galeria, md, flags=re.S)

    def link(m: re.Match) -> str:
        texto, alvo = m.group(1), m.group(2)
        if re.match(r"^([a-z]+:|#|mailto:)", alvo, flags=re.I):
            return m.group(0)
        caminho = (Path(base) / alvo).as_posix()
        partes = []
        for p in caminho.split("/"):
            if p == "..":
                partes and partes.pop()
            elif p not in ("", "."):
                partes.append(p)
        return f"[{texto}]({REPO_WEB}{'/'.join(partes)})"

    md = re.sub(r"(?<!!)\[([^\]]+)\]\(([^)\s]+)\)", link, md)
    # demais imagens com caminho relativo -> absoluto
    md = re.sub(r"!\[([^\]]*)\]\((?![a-z]+:|/)([^)\s]+)\)",
                lambda m: f"![{m.group(1)}]({(RAIZ / base / m.group(2)).resolve().as_posix()})", md)
    return md


def montar_markdown() -> str:
    partes = []
    for n, titulo, fonte, md, base in ENTREGAVEIS:
        corpo = converter(md, base)
        partes.append(f"# {n} · {titulo}\n\n*Fonte no repositório: {fonte}*\n\n{corpo}\n")
    return "\n".join(partes)


def referencia(destino: Path) -> None:
    """reference.docx do pandoc com a cara do Match: títulos em vermelho, cada entregável numa página nova."""
    dados = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], check=True,
                           capture_output=True).stdout
    destino.write_bytes(dados)
    doc = Document(str(destino))
    # o reference.docx do pandoc grava "Heading 1" com maiúscula e o doc.styles[...] do python-docx não acha
    st = {s.name: s for s in doc.styles}
    for nome in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if nome in st:
            st[nome].font.name = "Calibri"
            st[nome].font.size = Pt(10.5)
    for nome, tam in (("Heading 1", 22), ("Heading 2", 15), ("Heading 3", 12.5), ("Title", 28), ("Subtitle", 13)):
        s = st[nome]
        s.font.name = "Calibri"
        s.font.size = Pt(tam)
        s.font.bold = True
        s.font.color.rgb = VERMELHO if nome != "Subtitle" else RGBColor(0x55, 0x56, 0x5A)
    st["Heading 1"].paragraph_format.page_break_before = True
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(2)
    doc.save(str(destino))


def capa(doc: Document) -> None:
    corpo = doc.element.body
    primeiro = corpo[0]
    elementos = []

    def par(texto: str = "", tam: float = 11, cor: RGBColor | None = None, negrito: bool = False,
            centro: bool = True):
        p = doc.add_paragraph()
        if centro:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if texto:
            r = p.add_run(texto)
            r.font.size = Pt(tam)
            r.bold = negrito
            if cor:
                r.font.color.rgb = cor
        elementos.append(p._p)
        return p

    par("match", 40, VERMELHO, True)
    par(TITULO, 22, None, True)
    par(EVENTO, 12, RGBColor(0x55, 0x56, 0x5A))
    p = par()
    for f in TELAS_CAPA:
        if (RAIZ / f).exists():
            p.add_run().add_picture(str(RAIZ / f), width=Cm(4.6))
            p.add_run("   ")
    par("Entregáveis 01 a 06 do regulamento", 12, None, True)
    for n, titulo, *_ in ENTREGAVEIS:
        par(f"{n} · {titulo}", 11)
    par(f"Equipe: {EQUIPE}", 11)
    par(f"{DATA} · github.com/scorpion-bits/Hackaton-Ifood", 10, RGBColor(0x55, 0x56, 0x5A))
    par(AVISO, 9, RGBColor(0x8A, 0x8A, 0x8A))
    for e in elementos:
        primeiro.addprevious(e)


def tabela_galeria(doc: Document, paragrafo, itens) -> None:
    colunas = 4
    linhas = (len(itens) + colunas - 1) // colunas
    tab = doc.add_table(rows=linhas, cols=colunas)
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (img, titulo, texto) in enumerate(itens):
        cel = tab.cell(i // colunas, i % colunas)
        cel.width = Cm(4.25)
        p = cel.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if img.exists():
            p.add_run().add_picture(str(img), width=Cm(3.9))
        leg = cel.add_paragraph()
        r = leg.add_run(titulo + ". ")
        r.bold = True
        r.font.size = Pt(7.5)
        r2 = leg.add_run(texto)
        r2.font.size = Pt(7.5)
        r2.font.color.rgb = RGBColor(0x55, 0x56, 0x5A)
    paragrafo._p.addnext(tab._tbl)
    paragrafo._p.getparent().remove(paragrafo._p)


def bordas(tabela) -> None:
    el = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "DCDCDC")):
            e.set(qn(f"w:{k}"), v)
        el.append(e)
    tabela._tbl.tblPr.append(el)


def pos_processar(caminho: Path) -> None:
    doc = Document(str(caminho))
    for t in doc.tables:  # tabelas do Markdown (antes de entrarem as galerias, que ficam sem borda)
        bordas(t)
    for p in list(doc.paragraphs):
        m = re.fullmatch(r"GALERIA-(\d+)", p.text.strip())
        if m:
            tabela_galeria(doc, p, galerias[int(m.group(1))])
    largura_util, altura_max = Cm(17), Cm(22)
    for forma in doc.inline_shapes:
        for limite, medida in ((largura_util, "width"), (altura_max, "height")):
            atual = getattr(forma, medida)
            if atual > limite:
                k = limite / atual
                forma.width, forma.height = int(forma.width * k), int(forma.height * k)
    capa(doc)
    doc.core_properties.title = TITULO
    doc.core_properties.author = EQUIPE
    doc.core_properties.subject = "Entregáveis 01–06 · Hackathon SEMCOMP × iFood 2026"
    doc.save(str(caminho))


def main() -> int:
    if not shutil.which("pandoc"):
        print("ERRO: pandoc não encontrado. Instale em https://pandoc.org/installing.html", file=sys.stderr)
        return 1
    if not DIAGRAMAS.exists():
        print("AVISO: docs/entrega/diagramas/ não existe; rode `node gerar_pdf.mjs` antes para ter os diagramas.")
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    md = montar_markdown()
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        (tmp / "entregaveis.md").write_text(md, encoding="utf-8")
        ref = tmp / "referencia.docx"
        referencia(ref)
        subprocess.run(["pandoc", str(tmp / "entregaveis.md"), "-f", "gfm", "-t", "docx", "--toc", "--toc-depth=2",
                        "-M", "toc-title=Sumário", "--reference-doc", str(ref), "-o", str(SAIDA)],
                       check=True, cwd=str(AQUI))
    pos_processar(SAIDA)
    print(f"OK: {SAIDA.relative_to(RAIZ)} · {cont_diagrama} diagramas · {sum(map(len, galerias))} telas")
    return 0


if __name__ == "__main__":
    sys.exit(main())
